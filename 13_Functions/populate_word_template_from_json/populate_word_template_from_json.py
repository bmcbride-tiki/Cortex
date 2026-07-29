# =============================================================================
# populate_word_template_from_json.py
# -----------------------------------------------------------------------------
# Copyright 2025 Brian McBride at Tiki-1 Studio
# WHAT THIS FILE DOES
#   A Function: fills a .docx template's `{{key}}` placeholders from a
#   JSON object's keys, one placeholder per key (e.g. `{{name}}`,
#   `{{date}}`, `{{trade}}`). For a single `{{ content }}` token instead,
#   see `fill_docx_template.py`.
#
# WHAT IT INTERACTS WITH
#   - The given `template_path` .docx file, read via `python-docx`.
#   - `test_populate_word_template_from_json.py`, this file's paired test.
#   - `core_router.py`/`workflow_engine.py`, which dispatch this Function
#     the same way as any Task (generic `09_Functions` category), passing
#     its `data`/`template_path`/`output_dir` as a JSON payload.
# =============================================================================

import sys
import json
import time
from pathlib import Path
from typing import Dict, Any


class PopulateWordTemplateFromJson:
    """
    Fills a .docx template's {{key}} placeholders from a JSON object's keys, one
    placeholder per key (unlike fill_docx_template, which replaces a single {{ content }}
    token with plain text -- this is for templates with several distinct named fields,
    e.g. {{name}}, {{date}}, {{trade}}). Standalone-capable; also drag-and-droppable
    onto the Workflow Builder canvas (category 09_Functions), dispatched via CoreRouter
    same as any other Task.
    """

    def run(self, data: Dict[str, Any], template_path: str, output_dir: str) -> Dict[str, Any]:
        if not template_path:
            return {"success": False, "response": "A template_path is required."}
        if not isinstance(data, dict) or not data:
            return {"success": False, "response": "data must be a non-empty JSON object of {placeholder_key: value}."}

        from docx import Document

        src = Path(template_path)
        if not src.exists():
            return {"success": False, "response": f"Template file not found: {src}"}

        doc = Document(str(src))
        for para in doc.paragraphs:
            for key, value in data.items():
                for token in (f"{{{{ {key} }}}}", f"{{{{{key}}}}}"):
                    if token in para.text:
                        for run in para.runs:
                            run.text = run.text.replace(token, str(value))

        out_dir = Path(output_dir) if output_dir else src.parent
        out_dir.mkdir(parents=True, exist_ok=True)
        out_path = out_dir / f"{src.stem}_populated_{int(time.time())}.docx"
        doc.save(str(out_path))

        return {"success": True, "file_path": str(out_path)}


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(json.dumps({"success": False, "response": 'Expected one JSON payload arg: {"data": {"key": "value"}, "template_path": "...", "output_dir": "..."}'}))
        sys.exit(1)

    try:
        params = json.loads(sys.argv[1])
        result = PopulateWordTemplateFromJson().run(
            data=params.get("data", {}),
            template_path=params.get("template_path", ""),
            output_dir=params.get("output_dir", ""),
        )
    except Exception as e:
        result = {"success": False, "response": f"populate_word_template_from_json error: {e}"}

    print(json.dumps(result))
    if not result.get("success"):
        sys.exit(1)
